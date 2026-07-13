"""
Import service: parse word lists from Excel / JSON / TXT / clipboard, plus the
words-only lookup flow (fuzzy-match typed words against the dictionary).
"""
import difflib
import json


class ImportService:
    def __init__(self, word_repo, book_repo):
        self.word_repo = word_repo
        self.book_repo = book_repo

    # ── Lookup (words-only import) ────────────────────────────────────────────

    def lookup_words(self, raw_text: str) -> list:
        """
        Parse one-word-per-line text and match each against the dictionary.
        Returns [{input, matched, definition, example, chinese, found, method}]
        where method is 'exact' | 'fuzzy' | 'none'.
        """
        words = [w.strip().lower() for w in raw_text.splitlines() if w.strip()]
        results = []
        for word in words:
            row = self.word_repo.find_exact(word)
            if row:
                results.append(self._match(word, row, "exact"))
                continue
            rows = self.word_repo.find_like(word, limit=20)
            if rows:
                vocabs = [r[0] for r in rows]
                close = difflib.get_close_matches(word, vocabs, n=1, cutoff=0.6)
                row = rows[vocabs.index(close[0])] if close else rows[0]
                results.append(self._match(word, row, "fuzzy"))
                continue
            results.append({
                "input": word, "matched": None, "definition": "", "example": "",
                "chinese": "", "found": False, "method": "none",
            })
        return results

    @staticmethod
    def _match(word, row, method):
        return {
            "input": word, "matched": row[0],
            "definition": row[1] or "", "example": row[2] or "", "chinese": row[3] or "",
            "found": True, "method": method,
        }

    def import_word_matches(self, matches, bookname, new_book) -> str:
        if new_book and not self.book_repo.add_book(bookname):
            return f"Book '{bookname}' already exists!"
        words_dict = {
            m["matched"]: {
                "Definition": m.get("definition", ""),
                "ExampleSentence": m.get("example", ""),
                "Example_Chinese": m.get("chinese", ""),
            }
            for m in matches if m.get("found") and m.get("matched")
        }
        if not words_dict:
            return "No matching words to import."
        self.word_repo.add_words_to_book(bookname, words_dict, fetch_missing=False)
        return f"Successfully added {len(words_dict)} words!"

    # ── File / clipboard imports ──────────────────────────────────────────────

    def import_excel(self, file_path, sheet_name, word_column, definition_column,
                     example_sentence_column, example_chinese_column,
                     wordbook_name, new_book=False) -> str:
        import openpyxl

        def _cell(v):
            if v is None:
                return ""
            s = str(v).strip()
            return "" if s.lower() in ("none", "nan", "") else s

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            header = [_cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]

            if not word_column:
                return "Vocabulary column must be provided."
            required = [word_column] + [c for c in
                        (definition_column, example_sentence_column, example_chinese_column) if c]
            missing = [c for c in required if c not in header]
            if missing:
                return f"Columns not found in Excel file: {', '.join(missing)}"

            idx = {name: i for i, name in enumerate(header)}
            result, seen = {}, set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                word = _cell(row[idx[word_column]])
                if not word or word in seen:
                    continue
                seen.add(word)
                data = {}
                if definition_column and _cell(row[idx[definition_column]]):
                    data["Definition"] = _cell(row[idx[definition_column]])
                if example_sentence_column and _cell(row[idx[example_sentence_column]]):
                    data["ExampleSentence"] = _cell(row[idx[example_sentence_column]])
                if example_chinese_column and _cell(row[idx[example_chinese_column]]):
                    data["Example_Chinese"] = _cell(row[idx[example_chinese_column]])
                if data:
                    result[word] = data

            if new_book and not self.book_repo.add_book(wordbook_name):
                return f"Wordbook '{wordbook_name}' already exist!"
            self.word_repo.add_words_to_book(wordbook_name, result)
            return f"Successfully added {len(result)} vocabularies!"
        except FileNotFoundError:
            raise FileNotFoundError(f"Excel file not found at: {file_path}")
        except Exception as e:
            raise Exception(f"Error processing Excel file: {str(e)}")

    def import_json(self, file_path, vocab_key, definition_key, example_sentence_key,
                    example_chinese_key, wordbook_name, new_book=False) -> str:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except FileNotFoundError:
            raise FileNotFoundError(f"File not found: {file_path}")
        except json.JSONDecodeError as e:
            raise ValueError(f"Invalid JSON format: {e}")

        vocab_dict = {}
        if isinstance(data, list):
            self._process_list(data, vocab_dict, vocab_key, definition_key,
                               example_sentence_key, example_chinese_key)
        elif isinstance(data, dict):
            self._process_dict(data, vocab_dict, vocab_key, definition_key,
                               example_sentence_key, example_chinese_key)
        else:
            return f"Unsupported JSON structure. Expected list or dict, got {type(data)}"

        if new_book and not self.book_repo.add_book(wordbook_name):
            return f"Wordbook '{wordbook_name}' already exist!"
        self.word_repo.add_words_to_book(wordbook_name, vocab_dict)
        return f"Successfully added {len(vocab_dict)} vocabularies!"

    def import_txt(self, file_path, field_separator, entry_separator, wordbook_name,
                   new_book=False, vocab_field_index=0, definition_field_index=1,
                   example_field_index=2, example_chinese_field_index=3, encoding="utf-8") -> str:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read()
        except UnicodeDecodeError:
            try:
                with open(file_path, "r", encoding="latin-1") as f:
                    content = f.read()
            except Exception as e:
                return f"Error reading TXT file; Unknown Encoding. {e}"
        except Exception as e:
            return f"Error reading your TXT file: {e}"

        vocab_dict = self._parse_entries(content, field_separator, entry_separator,
                                         vocab_field_index, definition_field_index,
                                         example_field_index, example_chinese_field_index)
        if new_book and not self.book_repo.add_book(wordbook_name):
            return f"Wordbook '{wordbook_name}' already exist!"
        self.word_repo.add_words_to_book(wordbook_name, vocab_dict)
        return f"Successfully added {len(vocab_dict)} vocabularies!"

    def import_clipboard(self, file_content, field_separator, entry_separator, wordbook_name,
                         new_book=False, vocab_field_index=0, definition_field_index=1,
                         example_field_index=2, example_chinese_field_index=3) -> str:
        vocab_dict = self._parse_entries(file_content, field_separator, entry_separator,
                                         vocab_field_index, definition_field_index,
                                         example_field_index, example_chinese_field_index)
        if new_book and not self.book_repo.add_book(wordbook_name):
            return f"Wordbook '{wordbook_name}' already exist!"
        self.word_repo.add_words_to_book(wordbook_name, vocab_dict)
        return f"Successfully added {len(vocab_dict)} vocabularies!"

    # ── Parse helpers ─────────────────────────────────────────────────────────

    @staticmethod
    def _parse_entries(content, field_sep, entry_sep, vocab_idx, def_idx, ex_idx, ex_cn_idx) -> dict:
        vocab_dict = {}
        for entry in content.split(entry_sep):
            entry = entry.strip()
            if not entry:
                continue
            fields = [f.strip() for f in entry.split(field_sep)]
            vocab = fields[vocab_idx] if vocab_idx < len(fields) else ""
            if not vocab or vocab in vocab_dict:
                continue
            data = {}
            if def_idx < len(fields) and fields[def_idx]:
                data["Definition"] = fields[def_idx]
            if ex_idx < len(fields) and fields[ex_idx]:
                data["ExampleSentence"] = fields[ex_idx]
            if ex_cn_idx < len(fields) and fields[ex_cn_idx]:
                data["Example_Chinese"] = fields[ex_cn_idx]
            vocab_dict[vocab] = data if data else {}
        return vocab_dict

    def _process_dict(self, data_dict, vocab_dict, vocab_key, def_key, ex_key, ex_cn_key):
        first_key = next(iter(data_dict), None)
        if first_key and isinstance(data_dict.get(first_key), dict):
            sample = data_dict[first_key]
            if def_key in sample or ex_key in sample or ex_cn_key in sample:
                for word, wd in data_dict.items():
                    remapped = {}
                    if def_key in wd:
                        remapped["Definition"] = str(wd[def_key]).strip()
                    if ex_key in wd:
                        remapped["ExampleSentence"] = str(wd[ex_key]).strip()
                    if ex_cn_key in wd:
                        remapped["Example_Chinese"] = str(wd[ex_cn_key]).strip()
                    if remapped:
                        vocab_dict[word] = remapped
                return

        if any(isinstance(v, list) for v in data_dict.values()):
            for v in data_dict.values():
                if isinstance(v, list):
                    self._process_list(v, vocab_dict, vocab_key, def_key, ex_key, ex_cn_key)
                    break
        else:
            vocab = data_dict.get(vocab_key)
            if vocab:
                data = {}
                if data_dict.get(def_key) is not None:
                    data["Definition"] = str(data_dict[def_key]).strip()
                if data_dict.get(ex_key) is not None:
                    data["ExampleSentence"] = str(data_dict[ex_key]).strip()
                if data_dict.get(ex_cn_key) is not None:
                    data["Example_Chinese"] = str(data_dict[ex_cn_key]).strip()
                if data:
                    vocab_dict[str(vocab).strip()] = data

    def _process_list(self, data_list, vocab_dict, vocab_key, def_key, ex_key, ex_cn_key):
        for item in data_list:
            vocab = item.get(vocab_key)
            if vocab is None:
                continue
            vocab = str(vocab).strip()
            if not vocab or vocab in vocab_dict:
                continue
            data = {}
            if item.get(def_key) is not None:
                data["Definition"] = str(item[def_key]).strip()
            if item.get(ex_key) is not None:
                data["ExampleSentence"] = str(item[ex_key]).strip()
            if item.get(ex_cn_key) is not None:
                data["Example_Chinese"] = str(item[ex_cn_key]).strip()
            if data:
                vocab_dict[vocab] = data
