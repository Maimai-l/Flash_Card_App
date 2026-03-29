import datetime
import json
import logging
import os
import warnings
from typing import Any, Dict, List, Optional

from data.db.connection import DatabaseConnection
from data.db.word_repo import WordRepository
from data.db.book_repo import BookRepository
from data.db.list_generator import ListGenerator
from paths import data_path, ensure_parent_dir, ensure_user_data_seeded


class LibraryService:
    def __init__(self):
        ensure_user_data_seeded()
        self.logger = logging.getLogger(__name__)
        self._database_dir = data_path("database")
        self.db_path = str(self._database_dir / "vocabulary.db")
        self._lists_dir = data_path("lists")
        self._fsrs_dir = data_path("fsrs")
        ensure_parent_dir(self._database_dir / "dummy")
        self._lists_dir.mkdir(parents=True, exist_ok=True)
        self._fsrs_dir.mkdir(parents=True, exist_ok=True)

        # Initialize database and run any pending migrations
        self.db_conn = DatabaseConnection(self.db_path)
        self.db_conn.initialize_database()
        self.word_repo = WordRepository(self.db_path)
        self.book_repo = BookRepository(self.db_path)
        self.list_gen = ListGenerator(self.db_path)

        self.renew_app_info()

    def _initialize_app_info(self):
        app_info_path = data_path("app_info.json")
        ensure_parent_dir(app_info_path)
        app_info = {
            "last_session_date": None,
            "daily_new_limit":   20,
            "selected_book":     None,
            "theme":             "BLUE",
        }
        with open(app_info_path, "w", encoding="UTF-8") as f:
            json.dump(app_info, f, indent=4)

    def renew_app_info(self):
        """Load app_info, migrating from old format if necessary."""
        app_info_path = data_path("app_info.json")
        if not app_info_path.exists():
            self._initialize_app_info()
            return

        try:
            with open(app_info_path, "r", encoding="UTF-8") as f:
                app_info = json.load(f)
        except (FileNotFoundError, json.JSONDecodeError):
            self._initialize_app_info()
            return

        # Migrate old format (has last_open_date) to new format
        if "last_open_date" in app_info or "last_learn_date" in app_info:
            new_info = {
                "last_session_date": None,
                "daily_new_limit":   app_info.get("list_length", app_info.get("daily_new_limit", 20)),
                "selected_book":     None,
                "theme":             app_info.get("theme", "BLUE"),
            }
            with open(app_info_path, "w", encoding="UTF-8") as f:
                json.dump(new_info, f, indent=4)

    def get_app_info(self):
        with open(data_path("app_info.json"), "r", encoding="UTF-8") as f:
            return json.load(f)

    def update_app_info(self, app_info):
        with open(data_path("app_info.json"), "w", encoding="UTF-8") as f:
            json.dump(app_info, f, indent=4)

    # ── Import methods ──────────────────────────────────────────

    def import_excel(
        self,
        file_path: str,
        sheet_name: str,
        word_column: str,
        definition_column: str,
        example_sentence_column: str,
        example_chinese_column: str,
        wordbook_name: str,
        new_book: bool = False,
    ) -> str:
        import openpyxl

        def _cell_val(v):
            """Return stripped string or empty string for None / empty cells."""
            if v is None:
                return ""
            s = str(v).strip()
            return "" if s.lower() in ("none", "nan", "") else s

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

            # Read header row
            header = [_cell_val(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]

            required_columns = []
            if not word_column:
                return "Vocabulary column must be provided."
            required_columns.append(word_column)
            for col in [definition_column, example_sentence_column, example_chinese_column]:
                if col:
                    required_columns.append(col)

            missing = [c for c in required_columns if c not in header]
            if missing:
                return f"Columns not found in Excel file: {', '.join(missing)}"

            col_idx = {name: i for i, name in enumerate(header)}

            result_dict = {}
            processed_words = set()

            for row in ws.iter_rows(min_row=2, values_only=True):
                word = _cell_val(row[col_idx[word_column]])
                if not word or word in processed_words:
                    continue
                processed_words.add(word)

                word_data = {}
                if definition_column:
                    val = _cell_val(row[col_idx[definition_column]])
                    if val:
                        word_data["Definition"] = val
                if example_sentence_column:
                    val = _cell_val(row[col_idx[example_sentence_column]])
                    if val:
                        word_data["ExampleSentence"] = val
                if example_chinese_column:
                    val = _cell_val(row[col_idx[example_chinese_column]])
                    if val:
                        word_data["Example_Chinese"] = val

                if word_data:
                    result_dict[word] = word_data

            if new_book:
                if not self.book_repo.add_book(wordbook_name):
                    return f"Wordbook '{wordbook_name}' already exist!"

            self.word_repo.add_words_to_book(wordbook_name, result_dict)
            return f"Successfully added {len(result_dict)} vocabularies!"

        except FileNotFoundError:
            raise FileNotFoundError(f"Excel file not found at: {file_path}")
        except Exception as e:
            raise Exception(f"Error processing Excel file: {str(e)}")

    def import_json(
        self,
        file_path: str,
        vocab_key: str,
        definition_key: str,
        example_sentence_key: str,
        example_chinese_key: str,
        wordbook_name: str,
        new_book: bool = False,
    ) -> str:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except FileNotFoundError:
            raise FileNotFoundError(f"File not found: {file_path}")
        except json.JSONDecodeError as e:
            raise ValueError(f"Invalid JSON format: {e}")

        vocab_dict = {}
        if isinstance(data, list):
            self._process_list_data(data, vocab_dict, vocab_key, definition_key, example_sentence_key, example_chinese_key)
        elif isinstance(data, dict):
            self._process_dict_data(data, vocab_dict, vocab_key, definition_key, example_sentence_key, example_chinese_key)
        else:
            return f"Unsupported JSON structure. Expected list or dict, got {type(data)}"

        if new_book:
            if not self.book_repo.add_book(wordbook_name):
                return f"Wordbook '{wordbook_name}' already exist!"
        self.word_repo.add_words_to_book(wordbook_name, vocab_dict)
        return f"Successfully added {len(vocab_dict)} vocabularies!"

    def import_txt(
        self,
        file_path: str,
        field_separator: str,
        entry_separator: str,
        wordbook_name: str,
        new_book: bool = False,
        vocab_field_index: int = 0,
        definition_field_index: int = 1,
        example_field_index: int = 2,
        example_chinese_field_index: int = 3,
        encoding: str = "utf-8",
    ) -> str:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read()
        except UnicodeDecodeError:
            try:
                with open(file_path, "r", encoding="latin-1") as f:
                    content = f.read()
            except Exception as e:
                return f"Error reading TXT file; Unknown Encoding. {e}"
        except Exception as e:
            return f"Error reading your TXT file: {e}"

        vocab_dict = self._parse_text_entries(
            content, field_separator, entry_separator,
            vocab_field_index, definition_field_index,
            example_field_index, example_chinese_field_index,
        )

        if new_book:
            if not self.book_repo.add_book(wordbook_name):
                return f"Wordbook '{wordbook_name}' already exist!"
        self.word_repo.add_words_to_book(wordbook_name, vocab_dict)
        return f"Successfully added {len(vocab_dict)} vocabularies!"

    def import_clipboard(
        self,
        file_content: str,
        field_separator: str,
        entry_separator: str,
        wordbook_name: str,
        new_book: bool = False,
        vocab_field_index: int = 0,
        definition_field_index: int = 1,
        example_field_index: int = 2,
        example_chinese_field_index: int = 3,
    ) -> str:
        vocab_dict = self._parse_text_entries(
            file_content, field_separator, entry_separator,
            vocab_field_index, definition_field_index,
            example_field_index, example_chinese_field_index,
        )

        if new_book:
            if not self.book_repo.add_book(wordbook_name):
                return f"Wordbook '{wordbook_name}' already exist!"
        self.word_repo.add_words_to_book(wordbook_name, vocab_dict)
        return f"Successfully added {len(vocab_dict)} vocabularies!"

    # ── List generation ─────────────────────────────────────────

    def generate_new_list_from_book(self, count: int, bookname: str = None, is_extra: bool = False):
        if is_extra:
            app_info = self.get_app_info()
            if app_info["can_generate_list"]:
                app_info["can_generate_list"] = False
            else:
                return
        self.list_gen.generate_vocab_list_with_fsrs(
            count=count,
            vocab_output_path=str(self._lists_dir),
            fsrs_output_path=str(self._fsrs_dir),
            book_name=bookname if bookname != "All" else None,
            is_extra=is_extra,
        )

    def generate_review_list(self):
        result = self.list_gen.generate_vocab_list_with_fsrs(
            count=100,
            vocab_output_path=str(self._lists_dir),
            fsrs_output_path=str(self._fsrs_dir),
            due_only=True,
        )
        if not result.get("success"):
            app_info = self.get_app_info()
            app_info["today_review_completed"] = True
            self.update_app_info(app_info)

    def load_back_fsrs_data(self, json_file_name):
        self.word_repo.import_fsrs_data(str(self._fsrs_dir / json_file_name))

    # ── Data access ─────────────────────────────────────────────

    def get_list_data(self, file_name: str) -> dict:
        with open(self._lists_dir / file_name, "r", encoding="UTF-8") as f:
            return json.load(f)

    def update_list_info(self, file_name: str, info: dict):
        list_path = self._lists_dir / file_name
        with open(list_path, "r", encoding="UTF-8") as f:
            data = json.load(f)
        data["XXInfoXX"] = info
        with open(list_path, "w", encoding="UTF-8") as f:
            json.dump(data, f, indent=4)

    def get_book_names(self, details: bool = False):
        return self.book_repo.get_all_books(details)

    def get_list_names(self, bookname: str) -> list[str]:
        result = []
        for name in os.listdir(self._lists_dir):
            with open(self._lists_dir / name, "r", encoding="UTF-8") as f:
                info = json.load(f)["XXInfoXX"]
                if info["From"] == bookname or bookname == "All":
                    result.append(name)
        return result

    def get_book_complete_percentage(self, bookname) -> dict:
        if bookname != "All":
            return self.book_repo.get_book_completion_stats(bookname)
        else:
            return self.book_repo.get_book_completion_stats()

    # ── FSRS session ─────────────────────────────────────────────

    def get_due_words_for_refresh(self, book_name: str, exclude_words: list = None) -> list:
        """
        Return words currently due (fsrs_state > 0, due_date <= now) that are
        NOT in the exclude list.  Called after every answer to detect words
        that have become due mid-session (e.g. Hard/Again in Learning state).
        """
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()

        book_filter = book_name and book_name != "All"
        exclude = [w.lower() for w in (exclude_words or [])]

        def _run(sql, params):
            cursor.execute(sql, params)
            return cursor.fetchall()

        if exclude:
            ph = ",".join("?" * len(exclude))
            excl_clause = f" AND lower(w.vocab) NOT IN ({ph})" if book_filter else f" AND lower(vocab) NOT IN ({ph})"
        else:
            excl_clause = ""

        if book_filter:
            rows = _run(f"""
                SELECT w.vocab, w.definition_zh, w.example_en, w.example_zh,
                       w.phone_us, w.phone_uk, w.stability, w.difficulty, w.due_date, w.fsrs_state
                FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ? AND w.fsrs_state > 0
                  AND date(w.due_date) <= date('now'){excl_clause}
                ORDER BY w.due_date ASC
            """, [book_name] + exclude)
        else:
            rows = _run(f"""
                SELECT vocab, definition_zh, example_en, example_zh,
                       phone_us, phone_uk, stability, difficulty, due_date, fsrs_state
                FROM Words WHERE fsrs_state > 0
                  AND date(due_date) <= date('now'){excl_clause}
                ORDER BY due_date ASC
            """, exclude)

        conn.close()

        def _to_dict(r):
            return {
                "word": r[0], "definition": r[1] or "",
                "example": r[2] or "", "chinese": r[3] or "",
                "phone_us": r[4] or "", "phone_uk": r[5] or "",
                "stability": r[6] or 0, "difficulty": r[7] or 0,
                "due_date": r[8], "state": r[9] or 0,
            }
        return [_to_dict(r) for r in rows]

    def get_session_words(self, book_name: str, new_limit: int = 20) -> list:
        """
        Return today's combined session word list driven entirely by FSRS.

        1. Due-for-review words  : fsrs_state > 0 AND due_date <= NOW()
           (sorted oldest-due first — these are highest priority)
        2. New words             : due_date IS NULL, random, up to new_limit
           (never been seen — enters the FSRS cycle for the first time)

        Both groups are returned together as one list.
        """
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()

        book_filter = book_name and book_name != "All"

        def _rows(sql, params=()):
            cursor.execute(sql, params)
            return cursor.fetchall()

        if book_filter:
            review_rows = _rows("""
                SELECT w.vocab, w.definition_zh, w.example_en, w.example_zh,
                       w.phone_us, w.phone_uk, w.stability, w.difficulty, w.due_date, w.fsrs_state
                FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ? AND w.fsrs_state > 0
                  AND date(w.due_date) <= date('now')
                ORDER BY w.due_date ASC
            """, (book_name,))

            new_rows = _rows("""
                SELECT w.vocab, w.definition_zh, w.example_en, w.example_zh,
                       w.phone_us, w.phone_uk, w.stability, w.difficulty, w.due_date, w.fsrs_state
                FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ? AND w.due_date IS NULL
                ORDER BY RANDOM() LIMIT ?
            """, (book_name, new_limit))
        else:
            review_rows = _rows("""
                SELECT vocab, definition_zh, example_en, example_zh,
                       phone_us, phone_uk, stability, difficulty, due_date, fsrs_state
                FROM Words WHERE fsrs_state > 0
                  AND date(due_date) <= date('now')
                ORDER BY due_date ASC
            """)

            new_rows = _rows("""
                SELECT vocab, definition_zh, example_en, example_zh,
                       phone_us, phone_uk, stability, difficulty, due_date, fsrs_state
                FROM Words WHERE due_date IS NULL
                ORDER BY RANDOM() LIMIT ?
            """, (new_limit,))

        conn.close()

        def _to_dict(r):
            return {
                "word": r[0], "definition": r[1] or "",
                "example": r[2] or "", "chinese": r[3] or "",
                "phone_us": r[4] or "", "phone_uk": r[5] or "",
                "stability": r[6] or 0, "difficulty": r[7] or 0,
                "due_date": r[8], "state": r[9] or 0,
            }

        # Review words come first (overdue priority), then new words
        return [_to_dict(r) for r in review_rows] + [_to_dict(r) for r in new_rows]

    def get_today_stats(self, book_name: str = None) -> dict:
        """
        Return today's due counts for the home page.
        Automatically checks what FSRS has scheduled — no manual trigger needed.
        """
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()

        book_filter = book_name and book_name != "All"

        app_info = self.get_app_info()
        new_limit = app_info.get("daily_new_limit", 20)

        if book_filter:
            cursor.execute("""
                SELECT COUNT(*) FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ? AND w.due_date IS NULL
            """, (book_name,))
        else:
            cursor.execute("SELECT COUNT(*) FROM Words WHERE due_date IS NULL")
        new_available = min(cursor.fetchone()[0], new_limit)

        if book_filter:
            cursor.execute("""
                SELECT COUNT(*) FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ? AND w.fsrs_state > 0
                  AND date(w.due_date) <= date('now')
            """, (book_name,))
        else:
            cursor.execute("""
                SELECT COUNT(*) FROM Words WHERE fsrs_state > 0
                  AND date(due_date) <= date('now')
            """)
        review_due = cursor.fetchone()[0]
        conn.close()

        today = str(datetime.date.today())
        return {
            "new_count":      new_available,
            "review_count":   review_due,
            "total_due":      new_available + review_due,
            "today_done":     app_info.get("last_session_date") == today,
            "daily_new_limit": new_limit,
            "selected_book":  app_info.get("selected_book"),
        }

    def complete_session(self):
        """Mark today's session as done and write calendar entry."""
        app_info = self.get_app_info()
        app_info["last_session_date"] = str(datetime.date.today())
        self.update_app_info(app_info)
        self.add_calender_info(1)

    def get_learned_words(self, book_name: str = None, limit: int = 20) -> list:
        """Return N random words that have been reviewed at least once (fsrs_state > 0)."""
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        book_filter = book_name and book_name != "All"
        if book_filter:
            cursor.execute("""
                SELECT w.vocab, w.definition_zh, w.example_en, w.example_zh,
                       w.phone_us, w.phone_uk, w.stability, w.difficulty, w.due_date, w.fsrs_state
                FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ? AND w.fsrs_state > 0
                ORDER BY RANDOM() LIMIT ?
            """, (book_name, limit))
        else:
            cursor.execute("""
                SELECT vocab, definition_zh, example_en, example_zh,
                       phone_us, phone_uk, stability, difficulty, due_date, fsrs_state
                FROM Words WHERE fsrs_state > 0 ORDER BY RANDOM() LIMIT ?
            """, (limit,))
        rows = cursor.fetchall()
        conn.close()
        def _to_dict(r):
            return {
                "word": r[0], "definition": r[1] or "",
                "example": r[2] or "", "chinese": r[3] or "",
                "phone_us": r[4] or "", "phone_uk": r[5] or "",
                "stability": r[6] or 0, "difficulty": r[7] or 0,
                "due_date": r[8], "state": r[9] or 0,
            }
        return [_to_dict(r) for r in rows]

    def get_debug_words(self, book_name: str = None, limit: int = 20) -> list:
        """Debug: return N random words from book, ignoring FSRS state."""
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        book_filter = book_name and book_name != "All"
        if book_filter:
            cursor.execute("""
                SELECT w.vocab, w.definition_zh, w.example_en, w.example_zh,
                       w.phone_us, w.phone_uk, w.stability, w.difficulty, w.due_date, w.fsrs_state
                FROM Words w
                JOIN Word_Book_Words wbw ON w.word_id = wbw.word_id
                JOIN Word_Book wb ON wbw.book_id = wb.book_id
                WHERE wb.book_name = ?
                ORDER BY RANDOM() LIMIT ?
            """, (book_name, limit))
        else:
            cursor.execute("""
                SELECT vocab, definition_zh, example_en, example_zh,
                       phone_us, phone_uk, stability, difficulty, due_date, fsrs_state
                FROM Words ORDER BY RANDOM() LIMIT ?
            """, (limit,))
        rows = cursor.fetchall()
        conn.close()
        def _to_dict(r):
            return {
                "word": r[0], "definition": r[1] or "",
                "example": r[2] or "", "chinese": r[3] or "",
                "phone_us": r[4] or "", "phone_uk": r[5] or "",
                "stability": r[6] or 0, "difficulty": r[7] or 0,
                "due_date": r[8], "state": r[9] or 0,
            }
        return [_to_dict(r) for r in rows]

    def get_db_stats(self) -> dict:
        """Debug: return schema version, word counts, FSRS state distribution."""
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()

        # Schema version
        try:
            cursor.execute("SELECT MAX(version) FROM schema_version")
            schema_ver = cursor.fetchone()[0] or 0
        except Exception:
            schema_ver = 0

        # Total words
        cursor.execute("SELECT COUNT(*) FROM Words")
        total_words = cursor.fetchone()[0]

        # Words with definition_zh
        cursor.execute("SELECT COUNT(*) FROM Words WHERE definition_zh IS NOT NULL AND definition_zh != ''")
        words_with_def = cursor.fetchone()[0]

        # FSRS state distribution
        cursor.execute("""
            SELECT fsrs_state, COUNT(*) FROM Words GROUP BY fsrs_state ORDER BY fsrs_state
        """)
        state_dist = {}
        state_labels = {0: "New", 1: "Learning", 2: "Review", 3: "Relearning"}
        for state, cnt in cursor.fetchall():
            label = state_labels.get(state, f"State{state}")
            state_dist[label] = cnt

        # Due now
        cursor.execute("""
            SELECT COUNT(*) FROM Words WHERE fsrs_state > 0 AND date(due_date) <= date('now')
        """)
        due_now = cursor.fetchone()[0]

        # New (never seen)
        cursor.execute("SELECT COUNT(*) FROM Words WHERE due_date IS NULL")
        never_seen = cursor.fetchone()[0]

        # Per-book counts
        cursor.execute("""
            SELECT wb.book_name, COUNT(wbw.word_id)
            FROM Word_Book wb
            LEFT JOIN Word_Book_Words wbw ON wb.book_id = wbw.book_id
            GROUP BY wb.book_id ORDER BY wb.book_name
        """)
        books = [{"name": r[0], "count": r[1]} for r in cursor.fetchall()]

        conn.close()
        return {
            "schema_version": schema_ver,
            "total_words": total_words,
            "words_with_definition": words_with_def,
            "due_now": due_now,
            "never_seen": never_seen,
            "fsrs_states": state_dist,
            "books": books,
        }

    # ── Data management ─────────────────────────────────────────

    def _delete_data(self, folder_path):
        try:
            if not os.path.exists(folder_path):
                return False
            deleted = 0
            for item in os.listdir(folder_path):
                item_path = os.path.join(folder_path, item)
                if os.path.isfile(item_path):
                    os.remove(item_path)
                    deleted += 1
                elif os.path.isdir(item_path):
                    self._delete_data(item_path)
            return True
        except Exception as e:
            self.logger.exception("Error cleaning folder %s: %s", folder_path, e)
            return False

    # Preset book names — created automatically on fresh install, never by the user
    PRESET_BOOK_NAMES = {'Pre_Generate', 'CET_4+6_edited', 'TOEFL', 'GRE_8000_Words'}

    def soft_reset_data(self):
        """
        Soft reset (user-facing):
          1. Resets ALL FSRS progress to zero (every word becomes 'new' again).
          2. Removes user-imported word books and words that only existed in those books.
          3. Keeps all preset books (CET_4+6_edited, TOEFL, GRE_8000_Words, Pre_Generate)
             and their word data intact.
          4. Resets last_session_date and calendar; preserves other settings.
        """
        import sqlite3
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        try:
            # 1. Reset every word's FSRS state to "never seen"
            cursor.execute("""
                UPDATE Words SET
                    stability    = 0,
                    difficulty   = 0,
                    due_date     = NULL,
                    last_review  = NULL,
                    fsrs_state   = 0,
                    fsrs_step    = 0,
                    updated_time = datetime('now')
            """)

            # 2. Find user-imported books (anything not in the preset set)
            ph = ','.join('?' * len(self.PRESET_BOOK_NAMES))
            cursor.execute(
                f"SELECT book_id FROM Word_Book WHERE book_name NOT IN ({ph})",
                list(self.PRESET_BOOK_NAMES)
            )
            user_book_ids = [r[0] for r in cursor.fetchall()]

            if user_book_ids:
                id_ph = ','.join('?' * len(user_book_ids))
                # 3. Remove word-book links for those books
                cursor.execute(
                    f"DELETE FROM Word_Book_Words WHERE book_id IN ({id_ph})",
                    user_book_ids
                )
                # 4. Remove the books themselves
                cursor.execute(
                    f"DELETE FROM Word_Book WHERE book_id IN ({id_ph})",
                    user_book_ids
                )
                # 5. Remove words that no longer belong to any book
                cursor.execute("""
                    DELETE FROM Words WHERE word_id NOT IN (
                        SELECT DISTINCT word_id FROM Word_Book_Words
                    )
                """)

            conn.commit()
        except Exception as e:
            conn.rollback()
            self.logger.exception("soft_reset_data failed: %s", e)
            raise
        finally:
            conn.close()

        # Reset session date but keep other preferences (theme, daily limit, etc.)
        app_info = self.get_app_info()
        app_info['last_session_date'] = None
        self.update_app_info(app_info)
        self.initialize_calender_info()

    def reset_data(self):
        """Full factory reset — wipes everything and rebuilds from scratch."""
        self._initialize_app_info()
        self._delete_data(self._database_dir)
        self._delete_data(self._fsrs_dir)
        self._delete_data(self._lists_dir)
        self.db_conn.initialize_database()
        self.initialize_calender_info()

    def initialize_calender_info(self):
        calender_path = data_path("calender_info.json")
        ensure_parent_dir(calender_path)
        with open(calender_path, "w", encoding="UTF-8") as f:
            json.dump("{}", f)

    def add_calender_info(self, value: int):
        calender_path = data_path("calender_info.json")
        with open(calender_path, "r", encoding="UTF-8") as f:
            data = json.load(f)
        data[datetime.datetime.today().isoformat().split("T")[0]] = value
        with open(calender_path, "w", encoding="UTF-8") as f:
            json.dump(data, f, indent=4)

    def get_calender_info(self) -> dict:
        with open(data_path("calender_info.json"), "r", encoding="UTF-8") as f:
            return json.load(f)

    # ── Private helpers ─────────────────────────────────────────

    def _parse_text_entries(
        self, content, field_sep, entry_sep,
        vocab_idx, def_idx, example_idx, example_cn_idx,
    ) -> dict:
        entries = content.split(entry_sep)
        vocab_dict = {}

        for entry in entries:
            entry = entry.strip()
            if not entry:
                continue
            fields = [f.strip() for f in entry.split(field_sep)]

            vocab = fields[vocab_idx] if vocab_idx < len(fields) else ""
            if not vocab or vocab in vocab_dict:
                continue

            word_data = {}
            if def_idx < len(fields) and fields[def_idx]:
                word_data["Definition"] = fields[def_idx]
            if example_idx < len(fields) and fields[example_idx]:
                word_data["ExampleSentence"] = fields[example_idx]
            if example_cn_idx < len(fields) and fields[example_cn_idx]:
                word_data["Example_Chinese"] = fields[example_cn_idx]

            vocab_dict[vocab] = word_data if word_data else {}

        return vocab_dict

    def _process_dict_data(self, data_dict, vocab_dict, vocab_key, def_key, example_key, example_cn_key):
        first_key = next(iter(data_dict), None)
        if first_key and isinstance(data_dict.get(first_key), dict):
            sample = data_dict[first_key]
            if def_key in sample or example_key in sample or example_cn_key in sample:
                for word, wd in data_dict.items():
                    remapped = {}
                    if def_key in wd:
                        remapped["Definition"] = str(wd[def_key]).strip()
                    if example_key in wd:
                        remapped["ExampleSentence"] = str(wd[example_key]).strip()
                    if example_cn_key in wd:
                        remapped["Example_Chinese"] = str(wd[example_cn_key]).strip()
                    if remapped:
                        vocab_dict[word] = remapped
                return

        if any(isinstance(v, list) for v in data_dict.values()):
            for v in data_dict.values():
                if isinstance(v, list):
                    self._process_list_data(v, vocab_dict, vocab_key, def_key, example_key, example_cn_key)
                    break
        else:
            vocab = data_dict.get(vocab_key)
            if vocab:
                vocab = str(vocab).strip()
                word_data = {}
                if data_dict.get(def_key) is not None:
                    word_data["Definition"] = str(data_dict[def_key]).strip()
                if data_dict.get(example_key) is not None:
                    word_data["ExampleSentence"] = str(data_dict[example_key]).strip()
                if data_dict.get(example_cn_key) is not None:
                    word_data["Example_Chinese"] = str(data_dict[example_cn_key]).strip()
                if word_data:
                    vocab_dict[vocab] = word_data

    def _process_list_data(self, data_list, vocab_dict, vocab_key, def_key, example_key, example_cn_key):
        for item in data_list:
            vocab = item.get(vocab_key)
            if vocab is None:
                continue
            vocab = str(vocab).strip()
            if not vocab or vocab in vocab_dict:
                continue

            word_data = {}
            if item.get(def_key) is not None:
                word_data["Definition"] = str(item[def_key]).strip()
            if item.get(example_key) is not None:
                word_data["ExampleSentence"] = str(item[example_key]).strip()
            if item.get(example_cn_key) is not None:
                word_data["Example_Chinese"] = str(item[example_cn_key]).strip()
            if word_data:
                vocab_dict[vocab] = word_data
